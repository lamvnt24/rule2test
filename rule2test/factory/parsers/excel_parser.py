"""Bounded XLSX template import. Formulas and merged data cells are rejected."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile,BadZipFile
from factory.exceptions import ConfigurationError
from .common import LocatedRow,ImportFailure,ImportIssue,document,read_input,MAX_ROWS
from .schema import MANIFEST_KEYS,SHEETS
from ._builder import compile_document

class ExcelParser:
    def __init__(self,header_mapping=None):
        self.mapping=header_mapping or {}
        if type(self.mapping) is not dict or not set(self.mapping)<=set(SHEETS):
            raise ConfigurationError("Header mapping must use Policies, RulesV1, RulesV2 or Tests")
        for sheet,mapping in self.mapping.items():
            if type(mapping) is not dict or not set(mapping)<=set(SHEETS[sheet]) or not all(type(x) is str and x.strip() for x in mapping.values()):
                raise ConfigurationError("Invalid header mapping")

    def parse_path(self,path):
        if Path(path).suffix.lower()!=".xlsx":raise ImportFailure((ImportIssue(Path(path).name,"Only .xlsx files are supported"),))
        return self.parse(read_input(path),Path(path).name)

    def parse(self,data,name="import.xlsx"):
        doc=document(data,name,"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        try:
            import defusedxml
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise ConfigurationError("Excel import requires: python -m pip install --user -r requirements-excel.txt") from exc
        try:
            with ZipFile(BytesIO(data)) as archive:
                parts=archive.infolist()
                if len(parts)>2000 or sum(p.file_size for p in parts)>50*1024*1024:
                    raise ValueError("XLSX archive exceeds 2000 parts or 50 MiB expanded size")
                names=[p.filename for p in parts]
                if len(set(names))!=len(names):raise ValueError("Duplicate ZIP entries")
                if any("vbaproject" in n.lower() or "/externallinks/" in n.lower() for n in names):
                    raise ValueError("Macros and external workbook links are unsupported")
            workbook=load_workbook(BytesIO(data),data_only=False,read_only=False,keep_links=False)
        except Exception as exc:
            raise ImportFailure((ImportIssue(name,"Cannot read XLSX: "+str(exc)),)) from exc
        try:
            issues=[];tables={};manifest=None
            required={"Manifest",*SHEETS}
            for sheet in sorted(required-set(workbook.sheetnames)):issues.append(ImportIssue(name,"Missing required sheet",sheet=sheet))
            for sheet in sorted(set(workbook.sheetnames)-required-{"Guide"}):issues.append(ImportIssue(name,"Unexpected sheet; move explanatory content to Guide",sheet=sheet))
            for title in sorted(required&set(workbook.sheetnames)):
                sheet=workbook[title]
                if sheet.max_row>MAX_ROWS+1 or sheet.max_column>64:
                    issues.append(ImportIssue(name,"Sheet exceeds 5000 data rows or 64 columns",sheet=title));continue
                if sheet.merged_cells.ranges:
                    for cell_range in sheet.merged_cells.ranges:issues.append(ImportIssue(name,"Merged cells are not allowed in data sheets",sheet=title,cell=str(cell_range)))
                    continue
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type=="f":issues.append(ImportIssue(name,"Formula cells are not allowed; supply reviewed literal values",sheet=title,cell=cell.coordinate))
                        elif cell.data_type=="e":issues.append(ImportIssue(name,"Excel error cell",sheet=title,cell=cell.coordinate))
                if title=="Manifest":
                    if sheet.cell(1,1).value!="key" or sheet.cell(1,2).value!="value":
                        issues.append(ImportIssue(name,"Manifest headers must be key,value",sheet=title,cell="A1:B1"));continue
                    if any(sheet.cell(1,c).value is not None for c in range(3,sheet.max_column+1)):
                        issues.append(ImportIssue(name,"Unexpected manifest header",sheet=title,cell="C1"))
                    values={};locations={}
                    for number in range(2,sheet.max_row+1):
                        key=sheet.cell(number,1).value;value=sheet.cell(number,2).value
                        if key is None and value is None:continue
                        if key not in MANIFEST_KEYS or key in values:
                            issues.append(ImportIssue(name,"Unknown/duplicate manifest key",sheet=title,cell="A"+str(number)));continue
                        values[key]=value;locations[key]="B"+str(number)
                        if any(sheet.cell(number,c).value is not None for c in range(3,sheet.max_column+1)):
                            issues.append(ImportIssue(name,"Unexpected manifest column data",sheet=title,cell="C"+str(number)))
                    for key in set(MANIFEST_KEYS)-set(values):issues.append(ImportIssue(name,"Missing manifest key: "+key,sheet=title,cell="A1"))
                    manifest=LocatedRow(values,doc,sheet=title,row_number=2,columns={k:"B" for k in MANIFEST_KEYS},locations=locations)
                    continue
                mapping=self.mapping.get(title,{})
                expected={mapping.get(k,k).strip().casefold():k for k in SHEETS[title]}
                if len(expected)!=len(SHEETS[title]):
                    issues.append(ImportIssue(name,"Header mapping contains duplicate labels",sheet=title,cell="A1"));continue
                columns={}
                for cell in sheet[1]:
                    if cell.value is None:
                        if any(sheet.cell(r,cell.column).value is not None for r in range(2,sheet.max_row+1)):
                            issues.append(ImportIssue(name,"Data column has no header",sheet=title,cell=cell.coordinate))
                        continue
                    key=expected.get(str(cell.value).strip().casefold())
                    if key is None or key in columns:
                        issues.append(ImportIssue(name,"Unknown/duplicate column header",sheet=title,cell=cell.coordinate));continue
                    columns[key]=cell.column
                for key in set(SHEETS[title])-set(columns):issues.append(ImportIssue(name,"Missing column: "+key,sheet=title,cell="A1"))
                if len(columns)!=len(SHEETS[title]):continue
                rows=[]
                for number in range(2,sheet.max_row+1):
                    values={key:sheet.cell(number,col).value for key,col in columns.items()}
                    if all(value is None for value in values.values()):continue
                    rows.append(LocatedRow(values,doc,sheet=title,row_number=number,columns={k:get_column_letter(v) for k,v in columns.items()}))
                tables[title]=rows
            if issues:raise ImportFailure(issues[:100])
            try:return compile_document(doc,manifest,tables)
            except ImportFailure:raise
            except (ValueError,TypeError,KeyError,ArithmeticError) as exc:
                raise ImportFailure((ImportIssue(name,str(exc)),)) from exc
        finally:workbook.close()
