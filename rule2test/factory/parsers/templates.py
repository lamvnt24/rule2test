"""Synthetic import fixtures and literal-only XLSX template writer."""
import json
from io import BytesIO
from .schema import SHEETS,JSON_KEYS,MANIFEST_KEYS

PROFILES=("eligibility","claim_review","deductible")

def demo_payload(profile="eligibility"):
    if profile not in PROFILES:raise ValueError("Unknown demo profile")
    rid="R-"+profile
    default={"eligibility":"deny","claim_review":"allow","deductible":"invalid"}[profile]
    payload=dict(schema_version=1,created_at="2020-01-01T00:00:00+00:00",created_by="Synthetic Rule2Test fixture",
        policies=[],rules_v1=[],rules_v2=[],tests=[])
    def rule(version,field,op,kind,value,outcome,quote,deductible=None):
        return dict(rule_id=rid,version=version,title=profile.replace("_"," ").title(),field=field,operator=op,
            value_type=kind,value=value,currency="VND" if kind=="money" else None,outcome=outcome,payout_amount=None,
            deductible=deductible,effective_from=None,effective_to=None,quote=quote)
    for version in (1,2):
        payload["policies"].append(dict(label="v"+str(version),table_id="DT-"+profile,version=version,hit_policy="unique",
            default_outcome=default,default_amount=None,default_deductible=None,currency=None,as_of=None))
        if profile=="eligibility":
            maximum=60 if version==1 else 65
            quote=f"Age from 18 through {maximum} inclusive is allowed; otherwise denied."
            rows=[rule(version,"age","ge","integer",18,"allow",quote),rule(version,"age","le","integer",maximum,"allow",quote)]
        elif profile=="claim_review":
            threshold="100000000" if version==1 else "150000000"
            rows=[rule(version,"claim_amount","gt","money",threshold,"review",f"Claims exceeding {threshold} VND require review; otherwise allowed.")]
        else:
            deductible="5000000" if version==1 else "10000000"
            rows=[rule(version,"claim_amount","ge","money","0","payout",f"Payout equals max(claim_amount - {deductible} VND, 0).",deductible)]
        payload["rules_v"+str(version)]=rows
    values={"eligibility":(18,60,61),"claim_review":(100000000,100000001),"deductible":(5000000,5000001)}[profile]
    for number in values:
        field="age" if profile=="eligibility" else "claim_amount"
        expected=("allow" if number<=60 else "deny") if profile=="eligibility" else (
            ("review" if number>100000000 else "allow") if profile=="claim_review" else "payout")
        payload["tests"].append(dict(test_id="EXISTING-"+str(number),revision=1,title=f"Existing {field} = {number}",
            inputs_json=[dict(field=field,kind="integer" if profile=="eligibility" else "money",
                value=number if profile=="eligibility" else str(number),currency=None if profile=="eligibility" else "VND")],
            expected_outcome=expected,expected_amount=str(max(number-5000000,0)) if profile=="deductible" else None,
            currency="VND" if profile=="deductible" else None,rule_ids=rid,kind="boundary",
            rationale="Synthetic baseline expected value specified against V1. Review the V2 proposal separately."))
    return payload

def json_bytes(payload):
    return (json.dumps(payload,ensure_ascii=False,indent=2,allow_nan=False)+"\n").encode("utf-8")

def xlsx_bytes(payload,header_mapping=None):
    # Explicit text typing also prevents strings beginning with '=' from becoming executable formulas.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    workbook=Workbook();workbook.remove(workbook.active)
    def write(sheet,row,values):
        for col,value in enumerate(values,1):
            cell=sheet.cell(row,col)
            if type(value) in (list,dict):value=json.dumps(value,ensure_ascii=False,separators=(",",":"))
            cell.value=value
            if type(value) is str:cell.data_type="s"
            cell.alignment=Alignment(vertical="top",wrap_text=True)
    manifest=workbook.create_sheet("Manifest");write(manifest,1,("key","value"))
    for row,key in enumerate(MANIFEST_KEYS,2):write(manifest,row,(key,payload[key]))
    for key,title in JSON_KEYS.items():
        sheet=workbook.create_sheet(title);columns=SHEETS[title]
        mapping=(header_mapping or {}).get(title,{})
        write(sheet,1,[mapping.get(column,column) for column in columns])
        for number,entry in enumerate(payload[key],2):write(sheet,number,[entry[column] for column in columns])
    guide=workbook.create_sheet("Guide")
    guidance=[
        "Rule2Test structured import template - synthetic data only",
        "Manifest: retain schema_version=1; created_at must include a timezone.",
        "Policies: exactly v1 and v2; unique or first hit policy. Supply as_of for effective-dated rules.",
        "RulesV1 / RulesV2: repeated rule_id rows mean AND; repeat title, version, action and dates.",
        "Supported business fields: age and claim_amount. Operators: eq, ne, lt, le, gt, ge, is_null, is_missing.",
        "Use plain decimal TEXT for fractional money; no separators or exponent. Currency must be uppercase, e.g. VND.",
        "Payout: supply exactly one of payout_amount or deductible. Deductible means max(claim_amount - deductible, 0).",
        "Tests: inputs_json contains field, kind, value, currency for each input; rule_ids reference V1 IDs.",
        "No formulas, merged data cells, macros, external links or unknown data columns. Blank rows are ignored.",
        "All populated rows, including hidden rows, are validated. One invalid row rejects the entire import.",
        "Japanese source quotes are supported. For Japanese headers provide an explicit header mapping.",
        "Import creates DRAFT. Analyze and open review before QA/SME decisions; importing does not approve tests.",
        "See docs/IMPORT_FORMAT.md for the complete contract and CLI commands."
    ]
    for row,value in enumerate(guidance,1):write(guide,row,(value,))
    for sheet in workbook:
        sheet.freeze_panes="A2"
        if sheet.title!="Guide":sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="16324F")
        for column in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(column)].width=24
    guide.column_dimensions["A"].width=120
    stream=BytesIO();workbook.save(stream);workbook.close()
    return stream.getvalue()
