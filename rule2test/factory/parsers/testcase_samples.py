"""Synthetic test-case sheets and rule sentences for "Load demo example". Fill-in data only."""
from io import BytesIO
from factory.providers.llm.synthetic import SOURCE_PAIRS

HEADERS=("Test Case ID","Title","Preconditions","Steps","Test Data","Expected Result")

# Row 5 of each sheet is deliberately unreadable so the confirmation step is part of the demo.
ROWS={
    "eligibility":(
        ("TC001","Khách hàng 60 tuổi","Hồ sơ hợp lệ","Gửi yêu cầu tham gia","Tuổi: 60","Được chấp nhận"),
        ("TC002","Khách hàng 61 tuổi","Hồ sơ hợp lệ","Gửi yêu cầu tham gia","Tuổi: 61","Bị từ chối"),
        ("TC003","Khách hàng 18 tuổi","Hồ sơ hợp lệ","Gửi yêu cầu tham gia","Tuổi: 18","Được chấp nhận"),
        ("TC004","Khách hàng 17 tuổi","Hồ sơ hợp lệ","Gửi yêu cầu tham gia","Tuổi: 17","Bị từ chối"),
        ("TC005","Khách hàng lớn tuổi","Hồ sơ hợp lệ","Gửi yêu cầu tham gia","Khách hàng cao tuổi","Tùy trường hợp"),
    ),
    "claim_review":(
        ("TC001","Yêu cầu đúng ngưỡng","Hợp đồng còn hiệu lực","Nộp yêu cầu bồi thường","Số tiền yêu cầu: 100.000.000 VND","Được chấp nhận"),
        ("TC002","Yêu cầu vượt ngưỡng 1 đồng","Hợp đồng còn hiệu lực","Nộp yêu cầu bồi thường","Số tiền yêu cầu: 100.000.001 VND","Chuyển xem xét thủ công"),
        ("TC003","Yêu cầu 150 triệu","Hợp đồng còn hiệu lực","Nộp yêu cầu bồi thường","Số tiền yêu cầu: 150.000.000 VND","Chuyển xem xét thủ công"),
        ("TC004","Yêu cầu nhỏ","Hợp đồng còn hiệu lực","Nộp yêu cầu bồi thường","Số tiền yêu cầu: 5.000.000 VND","Được chấp nhận"),
        ("TC005","Yêu cầu rất lớn","Hợp đồng còn hiệu lực","Nộp yêu cầu bồi thường","Số tiền lớn","Tùy thẩm định viên"),
    ),
    "deductible":(
        ("TC001","Yêu cầu bằng mức khấu trừ","Hợp đồng còn hiệu lực","Tính số tiền chi trả","Số tiền yêu cầu: 5.000.000 VND","Chi trả 0 VND"),
        ("TC002","Yêu cầu trên mức khấu trừ 1 đồng","Hợp đồng còn hiệu lực","Tính số tiền chi trả","Số tiền yêu cầu: 5.000.001 VND","Chi trả 1 VND"),
        ("TC003","Yêu cầu 20 triệu","Hợp đồng còn hiệu lực","Tính số tiền chi trả","Số tiền yêu cầu: 20.000.000 VND","Chi trả 15.000.000 VND"),
        ("TC004","Yêu cầu 0 đồng","Hợp đồng còn hiệu lực","Tính số tiền chi trả","Số tiền yêu cầu: 0 VND","Chi trả 0 VND"),
        ("TC005","Yêu cầu không rõ","Hợp đồng còn hiệu lực","Tính số tiền chi trả","Chưa có số liệu","Chi trả theo hợp đồng"),
    ),
}

# Plain sentences a business analyst would type. The pattern interpreter reads these offline.
RULE_SENTENCES={
    "eligibility":("Khách hàng từ 18 đến 60 tuổi được tham gia bảo hiểm.","Khách hàng từ 18 đến 65 tuổi được tham gia bảo hiểm."),
    "claim_review":("Yêu cầu bồi thường trên 100.000.000 VND phải được xem xét thủ công.","Yêu cầu bồi thường trên 150.000.000 VND phải được xem xét thủ công."),
    "deductible":("Số tiền chi trả bằng số tiền yêu cầu trừ đi mức khấu trừ 5.000.000 VND, tối thiểu 0 VND.",
                  "Số tiền chi trả bằng số tiền yêu cầu trừ đi mức khấu trừ 10.000.000 VND, tối thiểu 0 VND."),
}

def rule_examples():
    """Every example the rules page can load: Vietnamese sentences for the pattern reader, Japanese pairs for AI replay."""
    result=[]
    for profile,(current,new) in RULE_SENTENCES.items():
        result.append(dict(key=profile,label=profile.replace("_"," ").title()+" · Vietnamese sentences",current=current,new=new,engine="pattern"))
    for profile,(current,new) in SOURCE_PAIRS.items():
        result.append(dict(key=profile+"-ja",label=profile.replace("_"," ").title()+" · Japanese document (AI replay)",current=current,new=new,engine="provider"))
    return result

def csv_bytes(profile):
    import csv,io
    stream=io.StringIO()
    writer=csv.writer(stream,lineterminator="\n");writer.writerow(HEADERS);writer.writerows(ROWS[profile])
    return stream.getvalue().encode("utf-8-sig")

def xlsx_bytes(profile):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    workbook=Workbook();sheet=workbook.active;sheet.title="TestCases"
    sheet.append(HEADERS)
    for row in ROWS[profile]:sheet.append(row)
    for cell in sheet[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="16324F")
    for column in range(1,len(HEADERS)+1):sheet.column_dimensions[get_column_letter(column)].width=28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical="top",wrap_text=True)
    sheet.freeze_panes="A2"
    stream=BytesIO();workbook.save(stream);workbook.close()
    return stream.getvalue()

def sample_file(profile):
    """(filename, bytes). XLSX when openpyxl is installed, otherwise the same rows as CSV."""
    if profile not in ROWS:raise ValueError("Unknown sample profile")
    try:
        import openpyxl  # noqa: F401
        return "demo-testcases-"+profile+".xlsx",xlsx_bytes(profile)
    except ImportError:
        return "demo-testcases-"+profile+".csv",csv_bytes(profile)
