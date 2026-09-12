import copy,hashlib,importlib.util,json,unittest
from io import BytesIO
from datetime import date,datetime
from decimal import Decimal
from factory.parsers.json_parser import JsonParser
from factory.parsers.excel_parser import ExcelParser
from factory.parsers.common import ImportFailure
from factory.parsers.templates import demo_payload,json_bytes,xlsx_bytes,PROFILES
from factory.models import SourceReference,content_hash
from factory.models.test_case import TestCase
from factory.exceptions import ValidationError

EXCEL=importlib.util.find_spec("openpyxl") is not None and importlib.util.find_spec("defusedxml") is not None

class JsonImportTests(unittest.TestCase):
    def parse(self,payload):return JsonParser().parse(json_bytes(payload),"policy.json")
    def test_all_profiles_have_explicit_old_and_new_snapshots(self):
        for profile in PROFILES:
            with self.subTest(profile=profile):
                p=demo_payload(profile);bundle=self.parse(p)
                self.assertEqual(bundle.old_rules[0].version,1)
                self.assertEqual(bundle.new_rules[0].version,2)
                self.assertEqual(bundle.document.document_hash,hashlib.sha256(json_bytes(p)).hexdigest())
                self.assertEqual(bundle.new_rules[0].sources[0].json_pointer,"/rules_v2/0/quote")
                self.assertEqual(bundle.existing_tests[0].sources[0].json_pointer,"/tests/0/inputs_json")
    def test_repeat_import_preserves_snapshot_hashes(self):
        one=self.parse(demo_payload());two=self.parse(demo_payload())
        self.assertEqual(one,two);self.assertEqual(content_hash(one.new_rules[0]),content_hash(two.new_rules[0]))
    def test_japanese_quote_is_preserved_verbatim(self):
        p=demo_payload();p["rules_v2"][0]["quote"]="  加入年齢は18歳以上65歳以下です。  "
        self.assertEqual(self.parse(p).new_rules[0].sources[0].quote,p["rules_v2"][0]["quote"])
    def test_rejects_duplicate_keys_nonfinite_and_invalid_encoding(self):
        for data in (b'{"schema_version":1,"schema_version":1}',b'{"x":NaN}',b'\xff',b'[]',b''):
            with self.subTest(data=data),self.assertRaises(ImportFailure):JsonParser().parse(data)
    def test_unknown_and_missing_columns_are_located(self):
        p=demo_payload();p["rules_v1"][0]["unexpected"]=7;del p["tests"][0]["rationale"]
        with self.assertRaises(ImportFailure) as ctx:self.parse(p)
        pointers={i.pointer for i in ctx.exception.issues}
        self.assertIn("/rules_v1/0/unexpected",pointers);self.assertIn("/tests/0/rationale",pointers)
    def test_bad_rows_are_aggregated_without_partial_result(self):
        p=demo_payload();p["rules_v1"][0]["value"]=True;p["rules_v2"][1]["value"]=999
        with self.assertRaises(ImportFailure) as ctx:self.parse(p)
        pointers={i.pointer for i in ctx.exception.issues}
        self.assertIn("/rules_v1/0/value",pointers);self.assertIn("/rules_v2/1/value",pointers)
    def test_nested_error_preserves_precise_operator_location(self):
        p=demo_payload();p["rules_v2"][0]["operator"]="approximately"
        with self.assertRaises(ImportFailure) as ctx:self.parse(p)
        self.assertEqual(ctx.exception.issues[0].pointer,"/rules_v2/0/operator")
    def test_rule_group_conflicting_metadata_is_rejected(self):
        p=demo_payload();p["rules_v1"][1]["title"]="different"
        with self.assertRaises(ImportFailure) as ctx:self.parse(p)
        self.assertTrue(any(i.pointer=="/rules_v1/1/rule_id" for i in ctx.exception.issues))
    def test_unknown_test_field_and_rule_reference_are_rejected(self):
        for key in ("field","reference"):
            p=demo_payload()
            if key=="field":p["tests"][0]["inputs_json"][0]["field"]="income"
            else:p["tests"][0]["rule_ids"]="missing-rule"
            with self.subTest(key=key),self.assertRaises(ImportFailure):self.parse(p)
    def test_null_missing_and_invalid_business_test_values_are_preserved(self):
        for kind,value in (("null",None),("missing",None),("integer",-1),("text","unknown")):
            p=demo_payload();p["tests"][0]["inputs_json"][0].update(kind=kind,value=value)
            p["tests"][0]["expected_outcome"]="invalid"
            t=self.parse(p).existing_tests[0]
            self.assertEqual(t.inputs[0].value.kind.value,kind);self.assertEqual(t.inputs[0].value.data,value)
    def test_decimal_money_is_lossless(self):
        p=demo_payload("claim_review")
        for key in ("rules_v1","rules_v2"):p[key][0]["value"]="100000000.123456789012"
        self.assertEqual(self.parse(p).new_rules[0].conditions[0].value.data,Decimal("100000000.123456789012"))
    def test_money_separator_exponent_precision_and_mixed_currency_rejected(self):
        for value in ("1,000","1e3","0.1234567890123"):
            p=demo_payload("claim_review");p["rules_v2"][0]["value"]=value
            with self.subTest(value=value),self.assertRaises(ImportFailure):self.parse(p)
        p=demo_payload("deductible")
        p["policies"][1].update(default_outcome="payout",default_amount="0",currency="USD")
        with self.assertRaises(ImportFailure):self.parse(p)
    def test_effective_dates_require_policy_as_of(self):
        p=demo_payload()
        for row in p["rules_v2"]:row.update(effective_from="2025-01-01",effective_to="2026-12-31")
        with self.assertRaises(ImportFailure):self.parse(p)
        p["policies"][1]["as_of"]="2026-01-01"
        self.assertEqual(self.parse(p).new_as_of,date(2026,1,1))
    def test_duplicate_test_and_policy_are_rejected(self):
        for key in ("tests","policies"):
            p=demo_payload();p[key].append(copy.deepcopy(p[key][0]))
            with self.subTest(key=key),self.assertRaises(ImportFailure):self.parse(p)
    def test_schema_and_naive_timestamp_rejected(self):
        for key,value in (("schema_version",2),("created_at","2020-01-01T00:00:00")):
            p=demo_payload();p[key]=value
            with self.subTest(key=key),self.assertRaises(ImportFailure):self.parse(p)
    def test_legacy_source_and_test_serialization_unchanged(self):
        from tests.fixtures.engine_cases import approved_case
        from factory.models import Value,ValueKind,Action,Outcome
        _,rules,test,_=approved_case("eligibility",Value(kind=ValueKind.INTEGER,data=18),Action(outcome=Outcome.ALLOW))
        source=rules[0].sources[0]
        self.assertNotIn("json_pointer",source.to_dict());self.assertNotIn("sources",test.to_dict())
        self.assertEqual(SourceReference.from_json(source.to_json()).to_json(),source.to_json())
        self.assertEqual(TestCase.from_json(test.to_json()).to_json(),test.to_json())
    def test_unused_currency_is_rejected_instead_of_discarded(self):
        for place in ("policy","rule","test","input"):
            p=demo_payload()
            target={"policy":p["policies"][0],"rule":p["rules_v1"][0],"test":p["tests"][0],
                "input":p["tests"][0]["inputs_json"][0]}[place]
            target["currency"]="USD"
            with self.subTest(place=place),self.assertRaises(ImportFailure):self.parse(p)

    def test_json_pointer_validation(self):
        kwargs=dict(document_id="a",document_hash="0"*64,quote="x")
        with self.assertRaises(ValidationError):SourceReference(**kwargs,json_pointer="invalid")
        with self.assertRaises(ValidationError):SourceReference(**kwargs,json_pointer="/bad~2")
        with self.assertRaises(ValidationError):SourceReference(**kwargs,json_pointer="/x",sheet="S",cell="A1")

@unittest.skipUnless(EXCEL,"Install requirements-excel.txt for XLSX coverage")
class ExcelImportTests(unittest.TestCase):
    def workbook(self,profile="eligibility"):
        from openpyxl import load_workbook
        w=load_workbook(BytesIO(xlsx_bytes(demo_payload(profile))))
        self.addCleanup(w.close);return w
    def parse(self,w,mapping=None):
        data=BytesIO();w.save(data)
        return ExcelParser(mapping).parse(data.getvalue(),"policy.xlsx")
    def test_all_templates_match_json_business_semantics(self):
        for profile in PROFILES:
            with self.subTest(profile=profile):
                actual=self.parse(self.workbook(profile));expected=JsonParser().parse(json_bytes(demo_payload(profile)))
                self.assertEqual(actual.new_rules[0].conditions,expected.new_rules[0].conditions)
                self.assertEqual(actual.new_rules[0].action,expected.new_rules[0].action)
                self.assertEqual(actual.existing_tests[0].inputs,expected.existing_tests[0].inputs)
                self.assertEqual(actual.new_rules[0].sources[0].sheet,"RulesV2")
                self.assertEqual(actual.new_rules[0].sources[0].cell,"N2")
                self.assertIsNone(actual.new_rules[0].sources[0].json_pointer)
    def test_formula_error_and_merged_cells_rejected_at_location(self):
        for kind in ("formula","error","merged"):
            w=self.workbook()
            if kind=="formula":w["RulesV2"]["G2"]="=18+1"
            elif kind=="error":w["RulesV2"]["G2"]="#VALUE!"
            else:w["RulesV2"].merge_cells("A2:A3")
            with self.subTest(kind=kind),self.assertRaises(ImportFailure) as ctx:self.parse(w)
            self.assertTrue(any(i.sheet=="RulesV2" and i.cell==("A2:A3" if kind=="merged" else "G2") for i in ctx.exception.issues))
    def test_missing_sheet_duplicate_header_and_unknown_column(self):
        for kind in ("sheet","duplicate","extra"):
            w=self.workbook()
            if kind=="sheet":del w["Tests"]
            elif kind=="duplicate":w["RulesV2"]["B1"]="rule_id"
            else:w["RulesV2"]["O1"]="unsupported"
            with self.subTest(kind=kind),self.assertRaises(ImportFailure):self.parse(w)
    def test_japanese_header_mapping_and_quote(self):
        w=self.workbook();w["RulesV2"]["G1"]="閾値";w["RulesV2"]["N2"]="加入年齢は18歳以上です。"
        with self.assertRaises(ImportFailure):self.parse(w)
        b=self.parse(w,{"RulesV2":{"value":"閾値"}})
        self.assertEqual(b.new_rules[0].sources[0].quote,"加入年齢は18歳以上です。")
    def test_fractional_money_must_be_text(self):
        w=self.workbook("claim_review");w["RulesV2"]["G2"]=100.25
        with self.assertRaises(ImportFailure) as ctx:self.parse(w)
        self.assertTrue(any(i.cell=="G2" and i.sheet=="RulesV2" for i in ctx.exception.issues))
        w["RulesV2"]["G2"]="100.25"
        self.assertEqual(self.parse(w).new_rules[0].conditions[0].value.data,Decimal("100.25"))
    def test_native_date_cells_and_reordered_manifest(self):
        w=self.workbook()
        for row in (2,3):w["RulesV2"].cell(row,12,datetime(2025,1,1))
        w["Policies"]["I3"]=date(2026,1,1)
        self.assertEqual(self.parse(w).new_as_of,date(2026,1,1))
        w["Manifest"]["A2"]="created_at";w["Manifest"]["B2"]="bad-date"
        w["Manifest"]["A3"]="schema_version";w["Manifest"]["B3"]=1
        with self.assertRaises(ImportFailure) as ctx:self.parse(w)
        self.assertTrue(any(i.sheet=="Manifest" and i.cell=="B2" for i in ctx.exception.issues))
    def test_hidden_rows_are_validated(self):
        w=self.workbook();w["RulesV2"].row_dimensions[3].hidden=True;w["RulesV2"]["G3"]="not an integer"
        with self.assertRaises(ImportFailure) as ctx:self.parse(w)
        self.assertTrue(any(i.cell=="G3" for i in ctx.exception.issues))
    def test_blank_rows_are_ignored_but_partial_rows_rejected(self):
        w=self.workbook();w["RulesV2"]["A5"]="R-partial"
        with self.assertRaises(ImportFailure):self.parse(w)
        w["RulesV2"]["A5"]=None
        self.assertEqual(len(self.parse(w).new_rules),1)
    def test_macro_external_link_and_invalid_zip_rejected(self):
        from zipfile import ZipFile
        for entry in ("xl/vbaProject.bin","xl/externalLinks/externalLink1.xml"):
            data=BytesIO(xlsx_bytes(demo_payload()))
            with ZipFile(data,"a") as archive:archive.writestr(entry,b"test")
            with self.subTest(entry=entry),self.assertRaises(ImportFailure):ExcelParser().parse(data.getvalue())
        with self.assertRaises(ImportFailure):ExcelParser().parse(b"not a zip")
    def test_limit_rejects_far_out_cell(self):
        w=self.workbook();w["RulesV2"]["A5002"]="too many rows"
        with self.assertRaises(ImportFailure):self.parse(w)
